"""Preserve legacy Excel templates with installed Excel's native object model.

Legacy .xls charts and print layouts require Excel; templates are opened read-only.
Exports are written to a temporary staging folder, verified, then moved atomically.
"""
import json
import gc
import os
import shutil
import tempfile
import threading
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from analysis_engine import MONTHS, norm, sort_error, ERRORS

ROOT=Path(__file__).parent
TEMPLATE_ROOT=ROOT/'template_files'
EXCEL_LOCK=threading.Lock()
TEMPLATES={
 '202':('KABLO HURDA ANALİZİ.xls','KABLO HURDA PARETO.xls'),
 '204':('TERMİNAL HURDA ANALİZİ.xls','TERMİNAL HURDA PARETO.xls'),
 '206':('IDC HURDA ANALIZI.xlsx','IDC HURDA PARETO.xls'),
}

def desktop():
    import winreg
    with winreg.OpenKey(winreg.HKEY_CURRENT_USER,r'Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders') as key:
        path=winreg.QueryValueEx(key,'Desktop')[0]
    return Path(os.path.expandvars(path))

def write(sheet,row,col,matrix):
    if matrix:sheet.Range(sheet.Cells(row,col),sheet.Cells(row+len(matrix)-1,col+len(matrix[0])-1)).Value=tuple(tuple(r) for r in matrix)

def number_format(excel,range_,invariant):
    # Excel's COM dispatch follows the host UI locale on this Windows install.
    local=invariant.replace(',', '\x01').replace('.',excel.DecimalSeparator).replace('\x01',excel.ThousandsSeparator)
    range_.NumberFormatLocal=local

def clean_chart(chart):
    # Template annotations describe old months; logos are kept.
    for shape in list(chart.Shapes):
        if shape.Type not in (11,13):shape.Delete()

def literal(value):
    value=str(value or '')
    return "'"+value if value.startswith(('=','+','-','@')) else value

def new_series(chart,name,categories,values,kind,secondary=False):
    series=chart.SeriesCollection().NewSeries();series.Name=name
    series.XValues=categories;series.Values=values;series.ChartType=kind
    if secondary:series.AxisGroup=2
    if kind==4:series.MarkerStyle=8;series.MarkerSize=7
    return series

def clear_series(chart):
    while chart.SeriesCollection().Count:chart.SeriesCollection(1).Delete()

def analysis_book(excel,book,result,group):
    cable=group['group']=='202';width=11 if cable else 10
    for month in result['months']:
        if not any(norm(s.Name)==norm(MONTHS[month-1]) for s in book.Worksheets):
            book.Worksheets('OCAK').Copy(After=book.Worksheets(book.Worksheets.Count))
            book.Worksheets(book.Worksheets.Count).Name=MONTHS[month-1].upper()
    # Remove accidental duplicate NISAN tab; clear obsolete embedded input tables.
    seen=set()
    for sheet in list(book.Worksheets):
        month=next((m for m,label in enumerate(MONTHS,1) if norm(label)==norm(sheet.Name)),None)
        if month:
            if month in seen:sheet.Delete();continue
            seen.add(month)
            count=sheet.UsedRange.Row+sheet.UsedRange.Rows.Count-1
            sheet.Range(sheet.Cells(2,1),sheet.Cells(max(count,2),width)).ClearContents()
            items=[s for s in group['stocks'] if s['month']==month]
            if not items:
                if month in result['months']:sheet.Cells(2,1).Value='Analize uygun veri yok.'
                continue
            end=len(items)+2
            sheet.Range(sheet.Cells(2,1),sheet.Cells(2,width)).Copy(sheet.Range(sheet.Cells(2,1),sheet.Cells(end-1,width)))
            matrix=[]
            for i,item in enumerate(items,2):
                if cable:matrix.append([item['code'],item['unit'],item['consumption'],f'=C{i}*E{i}/1000',item['grams'],item['scrap'],f'=E{i}*F{i}/1000',item['unit_tl'],f'=F{i}*H{i}',item['unit_eur'],f'=F{i}*J{i}'])
                else:matrix.append([item['code'],'AD',item['consumption'],item['scrap'],None,None,item['unit_tl'],f'=D{i}*G{i}',item['unit_eur'],f'=D{i}*I{i}'])
            write(sheet,2,1,matrix)
            mixed=cable and any(i['unit']!='MT' for i in items)
            if mixed:
                sheet.Columns(2).Hidden=False
                sheet.Cells(1,3).Value='TÜKETİM (BİRİM)';sheet.Cells(1,6).Value='HURDA (BİRİM)';sheet.Cells(1,5).Value='GRAMAJ (g/birim)'
            sheet.Cells(end,1).Value='TOPLAM'
            total_cols=[3,4,6,7,9,11] if cable else [3,4,8,10]
            for col in total_cols:
                letter=chr(64+col);sheet.Cells(end,col).Formula=f'=SUM({letter}2:{letter}{end-1})'
            if mixed:
                sheet.Cells(end,3).ClearContents();sheet.Cells(end,6).ClearContents()
            sheet.Range(sheet.Cells(end,1),sheet.Cells(end,width)).Font.Bold=True
            quantity_cols=[3,6] if cable else [3,4]
            for col in quantity_cols:number_format(excel,sheet.Range(sheet.Cells(2,col),sheet.Cells(end,col)),'#,##0.####')
            for col in ([4,7] if cable else []):number_format(excel,sheet.Range(sheet.Cells(2,col),sheet.Cells(end,col)),'#,##0.000000')
            for col in ([8,10] if cable else [7,9]):number_format(excel,sheet.Range(sheet.Cells(2,col),sheet.Cells(end,col)),'0.000000')
            for col in ([9,11] if cable else [8,10]):number_format(excel,sheet.Range(sheet.Cells(2,col),sheet.Cells(end,col)),'#,##0.00')
            sheet.PageSetup.PrintArea=f'$A$1:${chr(64+width)}${end}'
            sheet.PageSetup.Zoom=False;sheet.PageSetup.FitToPagesWide=1;sheet.PageSetup.FitToPagesTall=False
        elif norm(sheet.Name).startswith('BIRIM MAL'):
            sheet.UsedRange.ClearContents()
            write(sheet,1,1,[['Stok Kodu','Ay','Birim Malzeme TL','Birim Malzeme EUR','Maliyet Tarihi']]+[[i['code'],i['month'],i['unit_tl'],i['unit_eur'],i['cost_date'][:10]] for i in group['stocks']])
    # Templates have old manually entered summaries; rebuild their owning data area.
    summary=book.Worksheets('GRAFİK VERİSİ');summary.UsedRange.ClearContents()
    write(summary,1,1,[[f"{group['name']} · {result['year']} · {', '.join(result['vouchers'])}"]])
    write(summary,3,1,[['Ay',f"Tüketim ({group['unit']})",f"Hurda ({group['unit']})",'PPM','Hedef','Hurda Maliyeti TL','Hurda Maliyeti EUR']])
    for i,m in enumerate(group['monthly'],4):
        month_sheet=next((s for s in book.Worksheets if norm(s.Name)==norm(m['label'])),None)
        items=[s for s in group['stocks'] if s['month']==m['month']]
        if items and month_sheet:
            end=len(items)+2;sn=month_sheet.Name.replace("'","''")
            consumption_col='D' if cable else 'C';qty_col='G' if cable else 'D';tl_col='I' if cable else 'H';eur_col='K' if cable else 'J'
            write(summary,i,1,[[m['label'],f"='{sn}'!{consumption_col}{end}",f"='{sn}'!{qty_col}{end}",f'=IF(B{i}>0,C{i}/B{i}*1000000,"")',group['target'],f"='{sn}'!{tl_col}{end}",f"='{sn}'!{eur_col}{end}"]])
        else:write(summary,i,1,[[m['label'],None,None,None,group['target'],None,None]])
    last=3+len(group['monthly']);total=last+2
    write(summary,total,1,[['DÖNEM TOPLAMI',f'=SUM(B4:B{last})',f'=SUM(C4:C{last})',f'=IF(B{total}>0,C{total}/B{total}*1000000,"")',group['target'],f'=SUM(F4:F{last})',f'=SUM(G4:G{last})']])
    if not group['stocks']:write(summary,total,1,[['DÖNEM TOPLAMI',None,None,None,group['target'],None,None]])
    number_format(excel,summary.Range(f'B4:G{total}'),'#,##0.00');summary.Columns('A:G').ColumnWidth=24
    summary.Range('A3:G3').Font.Bold=True
    summary.Cells(total+2,1).Value='PPM: analize dahil stokların tüketimi. Eksik verili stoklar hariç.'
    chart_sheet=next((s for s in book.Worksheets if s.ChartObjects().Count),None)
    if chart_sheet:
        chart=chart_sheet.ChartObjects(1).Chart;clear_series(chart)
    else:
        chart_sheet=book.Worksheets.Add(After=book.Worksheets(book.Worksheets.Count));chart_sheet.Name='GRAFİK HURDA'
        chart=chart_sheet.ChartObjects().Add(20,20,880,430).Chart
    clean_chart(chart)
    cats=summary.Range(f'A4:A{last}')
    new_series(chart,'Tüketim',cats,summary.Range(f'B4:B{last}'),51)
    new_series(chart,'PPM',cats,summary.Range(f'D4:D{last}'),4,True)
    if group['target'] is not None:new_series(chart,'Hedef PPM',cats,summary.Range(f'E4:E{last}'),4,True)
    chart.HasTitle=True;chart.ChartTitle.Text=f"{group['name']} hurda analizi — {result['year']}\nGerçekleşen: {group['ppm']:.2f} PPM" if group['ppm'] is not None else f"{group['name']} — analize uygun veri yok"
    chart.Axes(2,1).HasTitle=True;chart.Axes(2,1).AxisTitle.Text=f"Tüketim ({group['unit']})"
    chart.Axes(2,2).HasTitle=True;chart.Axes(2,2).AxisTitle.Text='PPM'
    for axis in (chart.Axes(2,1),chart.Axes(2,2)):
        axis.MinimumScale=0;axis.MaximumScaleIsAuto=True
        number_format(excel,axis.TickLabels,'#,##0.00')
    add_details(book,result,group)

def add_details(book,result,group):
    details=book.Worksheets.Add(After=book.Worksheets(book.Worksheets.Count));details.Name='Hurda Detayı'
    rows=[r for r in result['detail'] if r['code'].startswith(group['group'])]
    factors={s['code']:s['grams']/1000 if group['group']=='202' and group['pareto_unit']=='kg' else 1 for s in group['stocks']}
    matrix=[['Ay','Fiş Grubu','Stok Kodu','Hata Kodu','Miktar','Kaynak','Orijinal Hata Kodu','Mod Ataması','Kaynak Miktarı','Kaynak Birimi']]
    matrix += [[r['month'],r['voucher'],r['code'],int(r['error']) if str(r['error']).isdigit() else literal(r['error']),r['qty']*factors[r['code']],'Parçalanmış' if r['source']=='fragmented' else 'Ana Hurda',literal(r['original_error']), 'Evet' if r['assigned'] else 'Hayır',r['qty'],r['unit']] for r in rows]
    write(details,1,1,matrix);details.Columns('A:H').AutoFit()
    pivot_sheet=book.Worksheets.Add(After=book.Worksheets(book.Worksheets.Count));pivot_sheet.Name='Hata Kodu Pivot'
    if rows:
        cache=book.PivotCaches().Create(1,f"'Hurda Detayı'!R1C1:R{len(matrix)}C8")
        pivot=cache.CreatePivotTable(pivot_sheet.Range('A3'),f"HataPivot{group['group']}")
        pivot.PivotFields('Stok Kodu').Orientation=1
        pivot.PivotFields('Hata Kodu').Orientation=2
        pivot.PivotFields('Ay').Orientation=3
        pivot.PivotFields('Fiş Grubu').Orientation=3
        pivot.AddDataField(pivot.PivotFields('Miktar'),'Toplam Hurda',-4157)
        pivot_sheet.Cells(1,1).Value=f"Miktar birimi: {group['pivot_unit']}"
        pivot.PivotFields('Hata Kodu').AutoSort(1,'Hata Kodu')
        pivot_sheet.Columns.AutoFit()
    else:pivot_sheet.Cells(1,1).Value='Analize uygun hurda kaydı yok.'
    excluded=book.Worksheets.Add(After=book.Worksheets(book.Worksheets.Count));excluded.Name='Analiz Dışı Kayıtlar'
    write(excluded,1,1,[['Stok Kodu','Ay','Neden']]+[[i['code'],i['month'],i['reason']] for i in result['issues'] if i['code'].startswith(group['group'])]);excluded.Columns('A:C').AutoFit()

def pareto_book(excel,book,result,group):
    for sheet in book.Worksheets:
        month=next((m for m,label in enumerate(MONTHS,1) if norm(label)==norm(sheet.Name)),None)
        if not month:continue
        counts={};costs={}
        for r in result['detail']:
            if r['month']==month and r['code'].startswith(group['group']):
                item=next(i for i in group['stocks'] if i['code']==r['code'] and i['month']==month)
                qty=r['qty']*(item['grams']/1000 if group.get('pareto_unit')=='kg' else 1)
                counts[r['error']]=counts.get(r['error'],0)+qty
                rate=item['unit_tl']
                costs[r['error']]=costs.get(r['error'],0)+r['qty']*rate
        ordered=sorted(counts,key=lambda e:(-counts[e],sort_error(e)))
        if len(ordered)>13:
            sheet.Rows(f'22:{21+len(ordered)-13}').Insert()
        total_row=max(23,10+len(ordered));sheet.Range(f'B10:H{total_row}').ClearContents()
        # Clear template author and update the reporting context.
        sheet.Cells(11,10).MergeArea.ClearContents()
        sheet.Cells(5,9).Value=f"{result['year']} {MONTHS[month-1].upper()}"
        sheet.Cells(6,2).Value=f"ÖLÇÜLEN DEĞER: {group['name'].upper()} HURDA ({group.get('pareto_unit',group['qty_unit'])})"
        sheet.Cells(8,8).Value='MALZEME\nMALİYETİ TL'
        for i,error in enumerate(ordered,10):
            description=ERRORS.get(int(error),'Tanımsız hata kodu') if str(error).isdigit() else str(error)
            write(sheet,i,2,[[literal(error),literal(description),counts[error],f'=D{i}/$D${total_row}',f'=SUM($E$10:E{i})',i-9,costs[error]]])
        sheet.Cells(total_row,2).Value='TOPLAM'
        for col in ['D','E','H']:sheet.Range(f'{col}{total_row}').Formula=f'=SUM({col}10:{col}{total_row-1})'
        if not ordered:sheet.Cells(10,3).Value='Analize uygun hurda kaydı yok.'
        number_format(excel,sheet.Range(f'E10:F{total_row}'),'0.00%')
        number_format(excel,sheet.Range(f'H10:H{total_row}'),'#,##0.00')
        sheet.Cells(total_row,8).Font.Bold=True
        chart=sheet.ChartObjects(1).Chart if sheet.ChartObjects().Count else sheet.ChartObjects().Add(30,370,640,320).Chart
        clear_series(chart)
        clean_chart(chart)
        if ordered:
            end=9+len(ordered);cats=sheet.Range(f'C10:C{end}')
            new_series(chart,'Hurda payı',cats,sheet.Range(f'E10:E{end}'),51)
            new_series(chart,'Kümülatif pay',cats,sheet.Range(f'F10:F{end}'),4)
            chart.Axes(2).MinimumScale=0;chart.Axes(2).MaximumScale=1;number_format(excel,chart.Axes(2).TickLabels,'0%')
        chart.HasTitle=True;chart.ChartTitle.Text=f"{group['name']} — {MONTHS[month-1]} {result['year']} Pareto"

def export_excel(result,destination=None):
    """Keep Excel COM lifetime out of the web server's request threads."""
    with EXCEL_LOCK:
        process=subprocess.run([sys.executable,str(Path(__file__).resolve()),'--worker'],input=json.dumps({'result':result,'destination':str(destination) if destination else None},ensure_ascii=False),text=True,encoding='utf-8',capture_output=True,timeout=240,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0),env={**os.environ,'PYTHONUTF8':'1'})
        if process.returncode:
            message=process.stderr.strip().splitlines()[-1] if process.stderr.strip() else 'Microsoft Excel işlemi tamamlanamadı.'
            raise RuntimeError(message)
        return json.loads(process.stdout)

def _export_excel_native(result,destination=None):
    import pythoncom
    import win32com.client
    with EXCEL_LOCK:
        pythoncom.CoInitialize();excel=None
        dest=Path(destination) if destination else desktop()/'Çıktı'
        dest.mkdir(parents=True,exist_ok=True)
        final=dest/datetime.now().strftime('Analiz_%Y%m%d_%H%M%S_%f')
        stage=Path(tempfile.mkdtemp(prefix='.analiz_',dir=dest));paths=[]
        try:
            excel=win32com.client.DispatchEx('Excel.Application')
            excel.Visible=False;excel.DisplayAlerts=False;excel.AskToUpdateLinks=False;excel.AutomationSecurity=3
            for group in result['groups']:
                for index,template in enumerate(TEMPLATES[group['group']]):
                    source=TEMPLATE_ROOT/template
                    if not source.exists():raise RuntimeError(f'Çıktı şablonu bulunamadı: {source.name}')
                    book=excel.Workbooks.Open(str(source),UpdateLinks=0,ReadOnly=True)
                    try:
                        if index==0:analysis_book(excel,book,result,group)
                        else:pareto_book(excel,book,result,group)
                        excel.CalculateFullRebuild()
                        target=stage/(Path(template).stem+'.xlsx')
                        book.SaveAs(str(target),FileFormat=51)
                        paths.append(target.name)
                    finally:
                        book.Close(SaveChanges=False)
                        book=None
            verify_exports(stage,result)
            stage.rename(final)
            return {'folder':str(final),'files':[str(final/p) for p in paths]}
        except Exception:
            shutil.rmtree(stage,ignore_errors=True)
            raise
        finally:
            gc.collect()
            if excel:excel.Quit()
            excel=None
            gc.collect()
            pythoncom.CoUninitialize()

def verify_exports(folder,result):
    import openpyxl
    for group in result['groups']:
        for index,template in enumerate(TEMPLATES[group['group']]):
            path=folder/(Path(template).stem+'.xlsx')
            wb=openpyxl.load_workbook(path,data_only=False)
            try:
                if not any(s._charts for s in wb):raise RuntimeError(f'{path.name}: gerçek Excel grafiği bulunamadı.')
                if index==0:
                    if group['stocks'] and not wb['Hata Kodu Pivot']._pivots:raise RuntimeError('Excel pivot tablosu oluşmadı.')
                    values=openpyxl.load_workbook(path,read_only=True,data_only=True)
                    try:
                        for month in result['months']:
                            sheet=next(s for s in values if norm(s.title)==norm(MONTHS[month-1]))
                            count=sum(s['month']==month for s in group['stocks'])
                            if count:
                                total=count+2;col=9 if group['group']=='202' else 8
                                expected=sum(s['cost_tl'] for s in group['stocks'] if s['month']==month)
                                actual=sheet.cell(total,col).value
                                if actual is None or abs(actual-expected)>0.001:raise RuntimeError(f'{path.name}: maliyet toplamı doğrulanamadı.')
                    finally:values.close()
            finally:wb.close()

if __name__=='__main__':
    payload=json.load(sys.stdin)
    print(json.dumps(_export_excel_native(payload['result'],payload.get('destination')),ensure_ascii=False))
