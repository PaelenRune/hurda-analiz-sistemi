import io
import time
import zipfile
import openpyxl
import pytest
import sys
import importlib.util
from werkzeug.security import generate_password_hash
from analysis_engine import analyze, parse_file
from test_analysis import fixture, opts
from web_excel import export_zip

def test_web_xlsx_cached_results_native_charts_and_pivot():
    result=analyze(fixture(),opts())
    with zipfile.ZipFile(io.BytesIO(export_zip(result))) as z:
        assert len(z.namelist())==6
        for name in z.namelist():
            raw=z.read(name)
            with zipfile.ZipFile(io.BytesIO(raw)) as parts:
                assert not any('externalLinks' in p for p in parts.namelist())
            wb=openpyxl.load_workbook(io.BytesIO(raw),data_only=False)
            cached=openpyxl.load_workbook(io.BytesIO(raw),data_only=True)
            if 'PARETO' in name:
                assert wb['EYLÜL']._charts
                assert cached['EYLÜL']['D2'].value==1
            else:
                assert wb['Hata Kodu Pivot']._pivots
                pivot=wb['Hata Kodu Pivot']._pivots[0]
                assert pivot.cache.recordCount>0
                assert len(pivot.cache.records.r)==pivot.cache.recordCount
                assert wb['GRAFİK VERİSİ']._charts
                assert cached['GRAFİK VERİSİ']['F2'].value>0
                assert cached['EYLÜL']['I2' if name.startswith('KABLO') else 'F2'].value>0
            wb.close();cached.close()

def test_memory_parser_and_multipart_never_spool(monkeypatch):
    from app import app
    import tempfile
    def forbidden(*args,**kwargs):raise AssertionError('Upload must never touch disk')
    monkeypatch.setattr(tempfile,'TemporaryFile',forbidden)
    monkeypatch.setattr(tempfile,'NamedTemporaryFile',forbidden)
    book=openpyxl.Workbook();s=book.active
    s.append(['Stok Kodu','Gramaj'])
    for i in range(70000):s.append([f'202{i:06}',4])
    data=io.BytesIO();book.save(data);data.seek(0)
    assert len(data.getvalue())>512000  # Cross Werkzeug's normal disk-spooling threshold.
    assert len(parse_file(data,'weights',filename='weights.xlsx')['records'])==70000
    data.seek(0)
    with app.test_client() as client:
        response=client.post('/api/upload/weights',data={'file':(data,'weights.xlsx')})
        assert response.status_code==200
        assert response.json['count']==70000

def test_web_auth_isolation_csrf_logout_and_no_cache(monkeypatch):
    from app import app, STATE, LOGIN_ATTEMPTS
    monkeypatch.setitem(app.config,'WEB_MODE',True)
    monkeypatch.setitem(app.config,'PASSWORD_HASH',generate_password_hash('test-password-123'))
    origin={'Origin':'https://example.com'}
    LOGIN_ATTEMPTS.clear()
    a=app.test_client();b=app.test_client()
    assert a.get('/',base_url='https://example.com').status_code==302
    assert a.get('/health',base_url='https://example.com').status_code==200
    assert a.get('/api/state',base_url='https://example.com').status_code==401
    page=a.get('/login',base_url='https://example.com')
    assert b'name="csrf_token"' in page.data
    with a.session_transaction(base_url='https://example.com') as s:login_token=s['csrf_token']
    assert a.post('/login',base_url='https://example.com',data={'password':'test-password-123'},headers=origin).status_code==403
    b.get('/login',base_url='https://example.com')
    with b.session_transaction(base_url='https://example.com') as s:foreign_token=s['csrf_token']
    assert a.post('/login',base_url='https://example.com',data={'password':'test-password-123','csrf_token':foreign_token},headers=origin).status_code==403
    # Browser proxies can rewrite Origin; a session-bound token still verifies the form.
    assert a.post('/login',base_url='https://example.com',data={'password':'wrong','csrf_token':login_token},headers={'Origin':'https://browser-proxy.example'}).status_code==401
    login=a.post('/login',base_url='https://example.com',data={'password':'test-password-123','csrf_token':login_token},headers={'Origin':'https://browser-proxy.example'})
    assert login.status_code==302
    page=a.get('/',base_url='https://example.com')
    assert b'name="csrf-token"' in page.data
    a.get('/api/state',base_url='https://example.com')
    with a.session_transaction(base_url='https://example.com') as s:
        assert set(s)=={'authenticated','id','csrf_token'}
        token=s['csrf_token']
        assert token!=login_token
        STATE[s['id']]['inputs']=fixture()
    secured={'Origin':'https://browser-proxy.example','X-CSRF-Token':token}
    assert a.post('/api/analyze',base_url='https://example.com',json=opts(),headers={'Origin':'https://evil.example'}).status_code==403
    assert a.post('/api/analyze',base_url='https://example.com',json=opts()).status_code==403
    assert a.post('/api/analyze',base_url='https://example.com',json=opts(),headers={**origin,'X-CSRF-Token':foreign_token}).status_code==403
    assert a.post('/api/analyze',base_url='https://example.com',json=opts(),headers=secured).status_code==200
    export=a.post('/api/export/excel',base_url='https://example.com',headers=secured)
    assert export.status_code==200 and export.data.startswith(b'PK')
    assert 'no-store' in export.headers['Cache-Control']
    assert b.get('/api/state',base_url='https://example.com').status_code==401
    assert a.post('/api/reset',base_url='https://example.com',headers=secured).status_code==200
    assert a.get('/api/state',base_url='https://example.com').json['files']=={}
    assert a.post('/api/logout',base_url='https://example.com',headers=secured).status_code==200
    assert a.get('/api/state',base_url='https://example.com').status_code==401
    assert a.post('/api/reset',base_url='https://example.com',headers=secured).status_code==403

def test_idle_session_drops_input_data(monkeypatch):
    from app import app, STATE
    monkeypatch.setitem(app.config,'SESSION_TTL_SECONDS',1)
    a=app.test_client();a.get('/api/state')
    with a.session_transaction() as s:
        old_id=s['id'];STATE[old_id]['inputs']=fixture();STATE[old_id]['touched']=time.monotonic()-2
    assert a.get('/api/state').json['files']=={}

def test_web_wrong_password_rate_limit(monkeypatch):
    from app import app, LOGIN_ATTEMPTS
    monkeypatch.setitem(app.config,'WEB_MODE',True)
    monkeypatch.setitem(app.config,'PASSWORD_HASH',generate_password_hash('test-password-123'))
    LOGIN_ATTEMPTS.clear()
    c=app.test_client()
    c.get('/login')
    with c.session_transaction() as s:token=s['csrf_token']
    for _ in range(10):assert c.post('/login',data={'password':'wrong','csrf_token':token}).status_code==401
    assert c.post('/login',data={'password':'wrong','csrf_token':token}).status_code==429
    LOGIN_ATTEMPTS.clear()

@pytest.mark.skipif(sys.platform!='win32' or not importlib.util.find_spec('win32com'),reason='Native Excel validation requires Windows')
@pytest.mark.parametrize('mixed',[False,True])
def test_web_pivot_refreshes_in_microsoft_excel(tmp_path,mixed):
    import subprocess
    from test_analysis import row
    from test_deep_review import mixed_fixture
    f=mixed_fixture() if mixed else fixture()
    f['scrap']['records'].extend([row(qty=7,error='2'),row(qty=3,error='10'),row(qty=4,error='Belirsiz')])
    r=analyze(f,opts())
    with zipfile.ZipFile(io.BytesIO(export_zip(r))) as z:
        for name in z.namelist():
            if name.startswith('KABLO') and 'PARETO' not in name:
                path=tmp_path/'native-check.xlsx';path.write_bytes(z.read(name))
    code="""
import sys,gc,pythoncom,win32com.client
pythoncom.CoInitialize()
excel=win32com.client.DispatchEx('Excel.Application')
excel.Visible=False;excel.DisplayAlerts=False
try:
    book=excel.Workbooks.Open(sys.argv[1],UpdateLinks=0,ReadOnly=True,CorruptLoad=0)
    try:
        pivot=book.Worksheets('Hata Kodu Pivot').PivotTables(1)
        assert pivot.RefreshTable()
        from math import isclose
        mixed=sys.argv[2]=='True'
        assert isclose(pivot.GetPivotData('Toplam Hurda').Value,.176 if mixed else 39)
        assert isclose(pivot.GetPivotData('Toplam Hurda','Hata Kodu',2).Value,.028 if mixed else 7)
        assert isclose(pivot.GetPivotData('Toplam Hurda','Hata Kodu',10).Value,.012 if mixed else 3)
        assert book.Worksheets('Hata Kodu Pivot').Cells(2,2).Value==2
        assert book.Worksheets('Hata Kodu Pivot').Cells(2,3).Value==10
        assert book.Worksheets('GRAFİK VERİSİ').ChartObjects().Count==1
        del pivot
    finally:book.Close(False)
finally:excel.Quit()
del book,excel
gc.collect()
"""
    checked=subprocess.run([sys.executable,'-c',code,str(path),str(mixed)],capture_output=True,text=True,timeout=60)
    assert checked.returncode==0,checked.stdout+checked.stderr
